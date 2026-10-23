"""Constructores de fixtures sintéticos (basados en ejemplos reales anonimizados).

Ningún archivo real se versiona: los archivos se generan en memoria con openpyxl / xlwt /
strings CSV, replicando layouts, encabezados y rarezas observadas (espacios en encabezados,
U+200B, mojibake, fechas m/d, montos como texto contable, etc.).
"""

from __future__ import annotations

import io

import openpyxl
import pytest


def xlsx_bytes(hojas: dict[str, list[list]], merges: dict[str, list[str]] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        for fila in filas:
            ws.append(fila)
        for rango in (merges or {}).get(nombre, []):
            ws.merge_cells(rango)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xls_bytes(hojas: dict[str, list[list]]) -> bytes:
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    for nombre, filas in hojas.items():
        ws = wb.add_sheet(nombre)
        for i, fila in enumerate(filas):
            for j, v in enumerate(fila):
                if v is not None:
                    ws.write(i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------
# Detalle liquidación PYA (versión "estado de cuenta", con espacios en encabezados)
# ---------------------------------------------------------------------------------------------
ENC_LIQ_ESTADO = [
    "Numero de Pedido",
    "Sucursal",
    "Fecha de Pedido",
    "Monto bruto de la venta",
    " Descuento comercial neto otorgado x PeYa ",
    "Descuento otorgado por el local",
    " Desc. neto otorgado por Peya a usuarios ",
    "Descuentos neto por Retiro en el local",
    " Descuento por Comodín pago por el local ",
    "Envío a cargo del local",
    "Monto neto de la Venta ($)",
    "Descuentos PedidosYa a cobrar",
    "Monto comisionable Servicios PedidosYa",
    "Servicio Ventas PedidosYA (%)",
    "Servicio Ventas PedidoYa ($)",
    "Monto venta + envió cobrado por el local",
    "Cargo por Pedidos con Plus",
    "Descuento tarifa de servicio al cliente",
    "Pedido con Comodin",
    "PLUS",
    "Método de pago",
    "Cobrado por",
    "Método de entrega",
]
ENC_LIQ_COMISIONES = [
    "Número del pedido",
    "Local",
    "Fecha del pedido",
    "Total del pedido",
    "Descuento en productos subsidiado por PedidosYa",
    "Descuento en productos subsidiado por el local",
    "Descuento al usuario pagado por PedidosYa",
    "Descuento por retiro en el local pagado por PedidosYa",
    "Tarifa de descuentos fugaces",
    "Envío a cargo del local",
    "Total pagado por el usuario",
    "Descuentos PedidosYa a cobrar",
    "Monto comisionable",
    "Porcentaje de comisión",
    "Comisión por pedido",
    "Monto venta + envió cobrado por el local",
    "Cargo por Pedidos con Plus",
    "Descuento tarifa de servicio al cliente",
    "Pedido con Comodin",
    "PLUS",
    "Método de pago",
    "Cobrado por",
    "Método de entrega",
]


def fila_liq(
    pedido,
    sucursal,
    fecha,
    bruto,
    comercial=0,
    local=0,
    usuarios=0,
    retiro=0,
    fugaz=0,
    neto=0,
    comisionable=0,
    comision=0,
    plus=0,
    metodo="Pago en la aplicación",
    entrega="Envío de PedidosYa",
):
    return [
        pedido,
        sucursal,
        fecha,
        bruto,
        comercial,
        local,
        usuarios,
        retiro,
        fugaz,
        0,
        neto,
        0,
        comisionable,
        "16%",
        comision,
        0,
        plus,
        0,
        "No",
        1 if plus else 0,
        metodo,
        "PedidosYa",
        entrega,
    ]


# Ejemplos reales
VENTA_2299350728 = fila_liq(
    2299350728,
    "Melt Pizzas - Quilpue",
    "9/26/2026",
    26370,
    local=6400,
    neto=19970,
    comisionable=19970,
    comision=3195,
)
VENTA_2297054261 = fila_liq(
    2297054261,
    "Melt Pizzas - Concón",
    "9/24/2026",
    16390,
    local=5051,
    usuarios="2.359",
    fugaz="3.144",
    neto=8532,
    comisionable=8195,
    comision=1311,
    plus=690,
)


@pytest.fixture
def liq_xlsx() -> bytes:
    ventas = [
        ENC_LIQ_ESTADO,
        VENTA_2299350728,
        VENTA_2297054261,
        fila_liq(
            2318734706,
            "Melt Pizzas - Quilpue",
            "9/22/2026",
            10000,
            neto=10000,
            comisionable=10000,
            comision=1600,
        ),
        fila_liq(
            2318734706,
            "Melt Pizzas - Quilpue",
            "9/22/2026",
            10000,
            neto=10000,
            comisionable=10000,
            comision=1600,
        ),
        ["Total", None, None, 999999],
    ]
    cargos = [ENC_LIQ_ESTADO, fila_liq(2297054261, "Melt Pizzas - Concón", "9/24/2026", 0, neto=-1500)]
    reintegros = [ENC_LIQ_ESTADO]
    return xlsx_bytes({"Ventas": ventas, "Cargos": cargos, "Reintegros": reintegros})


# ---------------------------------------------------------------------------------------------
# Texto simulado del PDF (semana 05/10/2026 - 11/10/2026)
# ---------------------------------------------------------------------------------------------
PDF_SEMANA_0510 = """Estado de cuenta semanal
Razón social: INVERSIONES VFR I SPA
RUT: 76.800.177-4
Total pagado por los usuarios $ 541.712.804
Total con pago en la aplicación $ 470.958.783
Total con pago fuera de la aplicación $ 70.754.021
Total con pago fuera de la aplicación a cobrar $ 0
Servicios PedidosYa -$ 88.439.523
Comisión por pedidos -$ 88.439.523
Cargos operativos -$ 13.466.315
Cargos por rechazar pedidos* -$ 159.961
Cargos por reclamos de los usuarios -$ 4.731.769
Costo por pedidos plus -$ 7.312.520
Costos de envio -$ 1.262.065
Publicidad dentro de la aplicación -$ 271.010
Tarifa de descuentos fugaces -$ 271.010
Reintegros $ 14.156.402
Descuentos a usuarios pagados por PedidosYa $ 8.832.741
Descuentos en productos subsidiados por PedidosYa $ 85.680
Descuentos por retiro en el local pagados por PedidosYa $ 609.121
Reintegros por pedidos cancelados por los usuarios* $ 4.628.861
IVA de descuentos no facturados -$ 1
Ajustes de liquidación $ 0
Impuestos -$ 16.673.939
Impuesto en comisiones - IVA -$ 16.673.939
Total liquidado del 05/10/26 al 11/10/26 $ 437.018.419
Total con pago fuera de la aplicación ya cobrado $ 0
El día 19/10/2026 procesaremos el pago de $ 437.018.419 en tu cuenta BANCO DE CREDITO E INVERSIONES terminada en 5421.
* Ver detalle en el archivo adjunto.
"""

PDF_SEMANA_2109 = """Estado de cuenta semanal
INVERSIONES VFR I SPA
Total pagado por los usuarios
$ 500.000.000
Servicios PedidosYa -$ 80.000.000
Cargos operativos -$ 10.000.000
Reintegros $ 12.000.000
Ajustes de liquidación $ 855.310
Cargos / Reintegros adicionales $ 855.310
Impuestos -$ 15.315.578
Total liquidado del 21/09/26 al 27/09/26 $ 407.539.732
El día 05/10/2026 procesaremos el pago en tu cuenta BANCO DE CREDITO E INVERSIONES terminada en 5421.
"""
